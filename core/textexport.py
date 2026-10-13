"""Saving the text of a PDF (for example a recognised scan) as Word, Excel or plain text, without extra libraries
for Word: the .docx is a small zip of XML files written here."""
from __future__ import annotations

import re
import zipfile
from dataclasses import dataclass, field
from pathlib import Path
from xml.sax.saxutils import escape


@dataclass
class Cell:
    text: str
    x0: float
    x1: float
    col: int | None = None       # column of a ruled table, counted from its left rule


@dataclass
class Row:
    """One printed line: its cells (separated by wide gaps), look and place, in points of the page."""
    cells: list[Cell]
    y0: float
    y1: float
    size: float
    bold: bool
    italic: bool
    color: str
    font: str
    x0: float = 0.0
    x1: float = 0.0
    table: int | None = None     # number of the ruled table this row is in

    @property
    def text(self) -> str:
        return "\t".join(c.text for c in self.cells)


@dataclass
class PageText:
    width: float
    height: float
    rows: list[Row] = field(default_factory=list)
    tables: list[list[float]] = field(default_factory=list)   # per ruled table: x of its vertical rules
    table_y: list[tuple[float, float]] = field(default_factory=list)


def _vertical_rules(page) -> list[tuple[float, float, float]]:
    """Long thin vertical lines of the page (x, y0, y1 in points), found in its picture."""
    import cv2
    import numpy as np
    pix = page.get_pixmap(dpi=100, alpha=False)
    img = np.frombuffer(pix.samples, np.uint8).reshape(pix.height, pix.width, pix.n)
    gray = img.mean(axis=2) if pix.n > 1 else img[:, :, 0]
    bw = (gray < 150).astype(np.uint8)
    k = max(20, int(0.035 * pix.height))
    vert = cv2.morphologyEx(bw, cv2.MORPH_OPEN, np.ones((k, 1), np.uint8))
    n, _, stats, _ = cv2.connectedComponentsWithStats(vert, connectivity=8)
    sc = page.rect.width / pix.width
    out = []
    for i in range(1, n):
        x, y, w, h, _a = (int(v) for v in stats[i])
        if w <= 8 and h >= k:
            out.append(((x + w / 2) * sc, y * sc, (y + h) * sc))
    return out


def _tables(rules: list[tuple[float, float, float]]) -> list[tuple[list[float], float, float]]:
    """Vertical rules that belong to one table: they overlap in height. Returns (rule xs, top, bottom)."""
    groups: list[list] = []
    for r in sorted(rules, key=lambda t: t[1]):
        for g in groups:
            if r[1] <= g[2] + 6 and r[2] >= g[1] - 6:
                g[0].append(r[0])
                g[1], g[2] = min(g[1], r[1]), max(g[2], r[2])
                break
        else:
            groups.append([[r[0]], r[1], r[2]])
    out = []
    for xs, y0, y1 in groups:
        xs.sort()
        merged: list[float] = []
        for x in xs:
            if merged and x - merged[-1] < 4:
                continue
            merged.append(x)
        if len(merged) >= 3:                          # two columns at least
            out.append((merged, y0, y1))
    return out


def extract(doc) -> list[PageText]:
    """Text lines of every page with their look. Invisible text of a scan counts: it is the recognised text."""
    pages: list[PageText] = []
    for page in doc:
        pt = PageText(page.rect.width, page.rect.height)
        try:
            tables = _tables(_vertical_rules(page))
        except Exception:
            tables = []
        pt.tables = [t[0] for t in tables]
        pt.table_y = [(t[1], t[2]) for t in tables]
        lines = []
        for block in page.get_text("dict").get("blocks", []):
            for ln in block.get("lines", []):
                spans = [s for s in ln.get("spans", []) if s.get("text", "").strip()]
                if not spans:
                    continue
                size = max(s["size"] for s in spans)
                lines.append((spans, size))
        # a row is everything that sits on one printed line, whatever block it came from
        lines.sort(key=lambda t: (t[0][0]["bbox"][1] + t[0][0]["bbox"][3]) / 2)
        groups: list[list] = []
        for spans, size in lines:
            cy = (spans[0]["bbox"][1] + spans[0]["bbox"][3]) / 2
            for g in groups:
                if abs(cy - g[0]) < 0.45 * size:
                    g[1].extend(spans)
                    g[2] = max(g[2], size)
                    break
            else:
                groups.append([cy, list(spans), size])
        for cy, spans, size in groups:
            spans.sort(key=lambda s: s["bbox"][0])
            cells: list[Cell] = []
            tid = next((i for i, (ty0, ty1) in enumerate(pt.table_y) if ty0 - 2 <= cy <= ty1 + 2), None)
            if tid is not None:
                xs = pt.tables[tid]
                for s in spans:
                    mid = (s["bbox"][0] + s["bbox"][2]) / 2
                    ci = next((k for k in range(len(xs) - 1) if xs[k] <= mid < xs[k + 1]), None)
                    if ci is None:
                        continue
                    t_ = s["text"].strip()
                    if cells and cells[-1].col == ci:
                        gap = s["bbox"][0] - cells[-1].x1
                        cells[-1].text += (" " if gap > 0.12 * size else "") + t_
                        cells[-1].x1 = s["bbox"][2]
                    else:
                        cells.append(Cell(t_, s["bbox"][0], s["bbox"][2], ci))
                if not cells:
                    tid = None
            if tid is None:
              for s in spans:
                x0, x1 = s["bbox"][0], s["bbox"][2]
                if cells and x0 - cells[-1].x1 <= 1.0 * size:       # a normal word gap: same cell
                    gap = x0 - cells[-1].x1
                    cells[-1].text += (" " if gap > 0.12 * size else "") + s["text"].strip()
                    cells[-1].x1 = x1
                else:
                    cells.append(Cell(s["text"].strip(), x0, x1))
            main = max(spans, key=lambda s: len(s["text"]))
            fname = main.get("font", "")
            flags = main.get("flags", 0)
            pt.rows.append(Row(
                cells=cells, y0=min(s["bbox"][1] for s in spans), y1=max(s["bbox"][3] for s in spans), size=size,
                bold=bool(flags & 16) or "bold" in fname.lower(), italic=bool(flags & 2) or "italic" in fname.lower(),
                color="%06X" % (main.get("color", 0) & 0xFFFFFF),
                font="Arial" if any(k in fname.lower() for k in ("arial", "helv", "sans", "calibri", "segoe", "verdana"))
                else "Times New Roman",
                x0=cells[0].x0, x1=cells[-1].x1, table=tid))
        pages.append(pt)
    return pages


# --------------------------------------------------------------------------- plain text

def save_txt(pages: list[PageText], path: Path):
    out = []
    for i, pg in enumerate(pages):
        if i:
            out.append("\f")
        out.extend(r.text.replace("\t", "    ") for r in pg.rows)
    Path(path).write_text("\n".join(out), encoding="utf-8")


# --------------------------------------------------------------------------- Word

def _paragraphs(pg: PageText) -> list[list[Row]]:
    """Lines that belong together (a wrapped paragraph) become one paragraph."""
    if not pg.rows:
        return []
    text_w = max((r.x1 - r.x0 for r in pg.rows), default=1.0)
    paras: list[list[Row]] = [[pg.rows[0]]]
    for r in pg.rows[1:]:
        p = paras[-1][-1]
        same = (abs(r.size - p.size) < 0.6 and r.bold == p.bold and len(r.cells) == 1 and len(p.cells) == 1
                and r.y0 - p.y1 < 0.55 * p.size and abs(r.x0 - paras[-1][0].x0) < 1.2 * p.size + 25 * (len(paras[-1]) == 1)
                and (p.x1 - p.x0) > 0.8 * text_w and not _centered(p, pg))
        if same:
            prev_t, next_t = p.text.rstrip(), r.text.lstrip()
            if re.match(r"\d+(\.\d+)*\.?\s+\S", next_t) or not (
                    next_t[:1].islower() or prev_t.endswith((",", "-", "—", "("))):
                same = False                  # only a lower-case or comma-led line goes on the same paragraph
        if same:
            paras[-1].append(r)
        else:
            paras.append([r])
    return paras


def _blocks(pg: PageText) -> list:
    """Paragraphs (lists of rows) and ruled tables ((\"table\", number, rows)) in reading order."""
    out: list = []
    buf: list[Row] = []

    def flush():
        if buf:
            out.extend(_paragraphs(PageText(pg.width, pg.height, list(buf))))
            buf.clear()

    i = 0
    while i < len(pg.rows):
        r = pg.rows[i]
        if r.table is None:
            buf.append(r)
            i += 1
            continue
        flush()
        j = i
        while j < len(pg.rows) and pg.rows[j].table == r.table:
            j += 1
        out.append(("table", r.table, pg.rows[i:j]))
        i = j
    flush()
    return out


def _is_number(t: str) -> bool:
    return bool(re.fullmatch(r"-?[\d\s\u00a0]+(?:[,.]\d+)?\s*[%]?", t.strip()))


def _table_xml(pg: PageText, tid: int, rows: list[Row], left: float) -> str:
    xs = pg.tables[tid]
    widths = [int((xs[k + 1] - xs[k]) * 20) for k in range(len(xs) - 1)]
    grid = "".join(f'<w:gridCol w:w="{w}"/>' for w in widths)
    edge = '<w:{0} w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
    borders = "".join(edge.format(n) for n in ("top", "left", "bottom", "right", "insideH", "insideV"))
    ind = int((xs[0] - left) * 20)
    out = [f'<w:tbl><w:tblPr><w:tblW w:w="{sum(widths)}" w:type="dxa"/><w:tblInd w:w="{max(0, ind)}" w:type="dxa"/>'
           f'<w:tblBorders>{borders}</w:tblBorders><w:tblLayout w:type="fixed"/>'
           '<w:tblCellMar><w:left w:w="60" w:type="dxa"/><w:right w:w="60" w:type="dxa"/></w:tblCellMar>'
           f'</w:tblPr><w:tblGrid>{grid}</w:tblGrid>']
    for r in rows:
        by_col = {c.col: c for c in r.cells}
        out.append("<w:tr>")
        for k, w in enumerate(widths):
            c = by_col.get(k)
            if c is None:
                para = "<w:p/>"
            else:
                jc = '<w:jc w:val="right"/>' if _is_number(c.text) and k > 0 else ""
                para = (f'<w:p><w:pPr><w:spacing w:before="20" w:after="20"/>{jc}</w:pPr>{_run(c.text, r)}</w:p>')
            out.append(f'<w:tc><w:tcPr><w:tcW w:w="{w}" w:type="dxa"/></w:tcPr>{para}</w:tc>')
        out.append("</w:tr>")
    out.append("</w:tbl>")
    return "".join(out)


def _centered(r: Row, pg: PageText) -> bool:
    return abs((r.x0 + r.x1) / 2 - pg.width / 2) < 0.03 * pg.width and r.x0 > 0.12 * pg.width and len(r.cells) == 1


def _run(text: str, r: Row) -> str:
    props = (f'<w:rFonts w:ascii="{r.font}" w:hAnsi="{r.font}" w:cs="{r.font}"/>'
             + ("<w:b/>" if r.bold else "") + ("<w:i/>" if r.italic else "")
             + (f'<w:color w:val="{r.color}"/>' if r.color not in ("000000", "") else "")
             + f'<w:sz w:val="{max(2, round(r.size * 2))}"/>')
    parts = text.split("\t")
    body = '<w:tab/>'.join(f'<w:t xml:space="preserve">{escape(p)}</w:t>' for p in parts)
    return f"<w:r><w:rPr>{props}</w:rPr>{body}</w:r>"


def save_docx(pages: list[PageText], path: Path):
    first = pages[0] if pages else PageText(595.0, 842.0)
    left = min((r.x0 for p in pages for r in p.rows), default=56.0)
    right = max((r.x1 for p in pages for r in p.rows), default=first.width - 56.0)
    top = min((r.y0 for p in pages for r in p.rows), default=56.0)
    margin_l = int(max(20.0, left) * 20)
    margin_r = int(max(20.0, first.width - right) * 20)
    margin_t = int(max(20.0, min(top, 120.0)) * 20)
    body = []
    for pi, pg in enumerate(pages):
        paras = _blocks(pg)
        prev_bottom = None
        for para in paras:
            if isinstance(para, tuple):                       # a ruled table
                body.append(_table_xml(pg, para[1], para[2], left))
                body.append('<w:p><w:pPr><w:spacing w:before="0" w:after="0" w:line="40" w:lineRule="exact"/></w:pPr></w:p>')
                prev_bottom = para[2][-1].y1
                continue
            head = para[0]
            text = " ".join(r.text for r in para)
            jc = '<w:jc w:val="center"/>' if _centered(head, pg) else ""
            gap = 0 if prev_bottom is None else max(0.0, head.y0 - prev_bottom)
            before = int(min(gap, 40.0) * 20)
            ind = ""
            if not jc and head.x0 - left > 8:
                ind = f'<w:ind w:left="{int((head.x0 - left) * 20)}"/>'
            if len(head.cells) > 1:       # columns: tab stops at the cell starts
                tabs = "".join(f'<w:tab w:val="left" w:pos="{int((c.x0 - left) * 20)}"/>' for c in head.cells[1:])
                ind += f"<w:tabs>{tabs}</w:tabs>"
            body.append(f'<w:p><w:pPr>{ind}<w:spacing w:before="{before}" w:after="0" w:line="{int(head.size * 1.2 * 20)}" '
                        f'w:lineRule="exact"/>{jc}</w:pPr>{_run(text, head)}</w:p>')
            prev_bottom = para[-1].y1
        if pi < len(pages) - 1:
            body.append('<w:p><w:r><w:br w:type="page"/></w:r></w:p>')
    sect = (f'<w:sectPr><w:pgSz w:w="{int(first.width * 20)}" w:h="{int(first.height * 20)}"/>'
            f'<w:pgMar w:top="{margin_t}" w:right="{margin_r}" w:bottom="{margin_t}" w:left="{margin_l}" '
            'w:header="0" w:footer="0" w:gutter="0"/></w:sectPr>')
    document = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
                '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
                f'<w:body>{"".join(body)}{sect}</w:body></w:document>')
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("[Content_Types].xml",
                   '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
                   '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
                   '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
                   '<Default Extension="xml" ContentType="application/xml"/>'
                   '<Override PartName="/word/document.xml" '
                   'ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/>'
                   '</Types>')
        z.writestr("_rels/.rels",
                   '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
                   '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
                   '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/'
                   'officeDocument" Target="word/document.xml"/></Relationships>')
        z.writestr("word/document.xml", document)


# --------------------------------------------------------------------------- Excel

_NUM = re.compile(r"^-?\d{1,3}(?:[  ]\d{3})*(?:[,.]\d+)?$|^-?\d{1,9}(?:[,.]\d+)?$")


def _value(text: str):
    """Short numbers and sums become numbers; long digit strings (INN, accounts, phone) stay text."""
    t = text.strip()
    if _NUM.match(t):
        digits = re.sub(r"\D", "", t)
        if len(digits) <= 9 or re.search(r"[,.]\d", t):
            try:
                v = float(t.replace(" ", "").replace(" ", "").replace(",", "."))
                return int(v) if v.is_integer() and not re.search(r"[,.]", t) else v
            except ValueError:
                pass
    return text


def _columns(rows: list[Row], tol: float) -> list[float]:
    xs = sorted(c.x0 for r in rows for c in r.cells)
    cols: list[list[float]] = []
    for x in xs:
        if cols and x - cols[-1][-1] <= tol:
            cols[-1].append(x)
        else:
            cols.append([x])
    return [sum(c) / len(c) for c in cols]


def save_xlsx(pages: list[PageText], path: Path):
    import openpyxl
    from openpyxl.styles import Alignment, Font
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for i, pg in enumerate(pages, 1):
        ws = wb.create_sheet(f"Страница {i}")
        if not pg.rows:
            continue
        tol = 1.2 * (sum(r.size for r in pg.rows) / len(pg.rows))
        multi = [r for r in pg.rows if len(r.cells) > 1 and r.table is None]
        cols = _columns(multi, tol) if multi else []
        widths: dict[int, int] = {}
        for ri, r in enumerate(pg.rows, 1):
            for c in r.cells:
                if r.table is not None and c.col is not None:
                    ci = c.col + 1                                # a ruled table: its own columns
                elif cols:
                    k = min(range(len(cols)), key=lambda k: abs(cols[k] - c.x0))
                    # a lone cell of a text line stays in the first column unless it stands under a table column
                    ci = k + 1 if len(r.cells) > 1 or abs(cols[k] - c.x0) <= tol else 1
                else:
                    ci = 1
                cell = ws.cell(row=ri, column=ci)
                if cell.value not in (None, ""):                 # two cells fell into one column: join
                    cell.value = f"{cell.value} {c.text}"
                else:
                    cell.value = _value(c.text)
                if r.bold:
                    cell.font = Font(bold=True)
                if r.table is not None and isinstance(cell.value, (int, float)):
                    cell.alignment = Alignment(horizontal="right")
                widths[ci] = max(widths.get(ci, 8), min(80, len(str(cell.value)) + 2))
        for ci, w in widths.items():
            ws.column_dimensions[openpyxl.utils.get_column_letter(ci)].width = w
    if not wb.sheetnames:
        wb.create_sheet("Страница 1")
    wb.save(path)


def save(doc, path) -> Path:
    """Write the text of ``doc`` to ``path``; the format follows the extension (.docx, .xlsx, .txt)."""
    path = Path(path)
    pages = extract(doc)
    ext = path.suffix.lower()
    if ext == ".docx":
        save_docx(pages, path)
    elif ext == ".xlsx":
        save_xlsx(pages, path)
    else:
        save_txt(pages, path)
    return path
