"""Turning any supported file into a list of RGB pages (PIL images).

Raster formats go through Pillow. PDF, SVG and text-like documents
(TXT, HTML, XLSX) are rendered with Qt, so no poppler/cairo is needed.
Qt rendering requires a QGuiApplication to exist (the GUI creates one,
tests create an offscreen one).
"""
from __future__ import annotations

import html
import threading
from pathlib import Path

from PIL import Image, ImageFile, ImageOps

ImageFile.LOAD_TRUNCATED_IMAGES = True
Image.MAX_IMAGE_PIXELS = None  # large scans and posters are legitimate input

try:
    import pillow_heif

    pillow_heif.register_heif_opener()
    HEIF_SUPPORTED = True
except ImportError:  # pragma: no cover - depends on the environment
    HEIF_SUPPORTED = False

RASTER_EXTS = {
    ".jpg", ".jpeg", ".jfif", ".png", ".webp", ".bmp", ".gif",
    ".tif", ".tiff", ".psd", ".heic", ".heif",
}
from .office import OFFICE_EXTS, OfficeError, can_convert, load_office

DOCUMENT_EXTS = {".pdf", ".svg", ".txt", ".html", ".htm"} | OFFICE_EXTS
ARCHIVE_EXTS = {".zip", ".rar", ".7z"}
SUPPORTED_EXTS = RASTER_EXTS | DOCUMENT_EXTS

A4_INCHES = (8.2677, 11.6929)
MAX_SIDE = 12000  # safety cap for vector renders
SHEET_ROW_LIMIT = 5000


_PDF_LOCK = threading.Lock()


class LoadError(Exception):
    """Raised when a file cannot be turned into images."""


def is_photo(path: Path) -> bool:
    """Raster images and SVG get crop presets; document pages do not."""
    return path.suffix.lower() in RASTER_EXTS or path.suffix.lower() == ".svg"


def to_rgb(img: Image.Image, background=(255, 255, 255)) -> Image.Image:
    """Flatten transparency onto white and return an RGB image.

    The embedded ICC profile (if any) is kept in ``info`` so it can be
    written to the output file.
    """
    # An RGB profile still describes the pixels after flattening; a gray or
    # CMYK one would not, so it is dropped.
    icc = img.info.get("icc_profile") if img.mode in ("RGB", "RGBA") else None
    if img.mode in ("RGBA", "LA", "PA") or (img.mode == "P" and "transparency" in img.info):
        rgba = img.convert("RGBA")
        base = Image.new("RGB", rgba.size, background)
        base.paste(rgba, mask=rgba.getchannel("A"))
        img = base
    elif img.mode == "CMYK":
        img = _cmyk_to_rgb(img)
    elif img.mode != "RGB":
        img = img.convert("RGB")
    img.info.pop("icc_profile", None)
    if icc:
        img.info["icc_profile"] = icc
    return img


def _cmyk_to_rgb(img: Image.Image) -> Image.Image:
    icc = img.info.get("icc_profile")
    if icc:
        try:
            import io
            from PIL import ImageCms

            src = ImageCms.ImageCmsProfile(io.BytesIO(icc))
            dst = ImageCms.createProfile("sRGB")
            return ImageCms.profileToProfile(img, src, dst, outputMode="RGB")
        except Exception:
            pass
    return img.convert("RGB")


# --------------------------------------------------------------------------- raster

def load_raster(path: Path, max_pages: int | None = None, draft_size: int | None = None) -> list[Image.Image]:
    try:
        img = Image.open(path)
    except Exception as exc:
        raise LoadError("файл повреждён или это не изображение") from exc

    if draft_size and img.format == "JPEG":
        img.draft("RGB", (draft_size, draft_size))

    pages: list[Image.Image] = []
    # Multi-page TIFFs become several pages; animated GIF/WebP use the first frame.
    n_frames = getattr(img, "n_frames", 1) if img.format == "TIFF" else 1
    for i in range(n_frames):
        if max_pages and i >= max_pages:
            break
        if n_frames > 1:
            img.seek(i)
        frame = ImageOps.exif_transpose(img) or img
        if draft_size:
            frame = frame.copy()
            frame.thumbnail((draft_size, draft_size), Image.Resampling.LANCZOS)
        page = to_rgb(frame)
        page.load()
        pages.append(page)
    return pages


# --------------------------------------------------------------------------- Qt helpers

def _qimage_to_pil(qimg) -> Image.Image:
    from PySide6.QtGui import QImage

    qimg = qimg.convertToFormat(QImage.Format.Format_RGBA8888)
    w, h = qimg.width(), qimg.height()
    buf = bytes(qimg.constBits())
    img = Image.frombuffer("RGBA", (w, h), buf, "raw", "RGBA", qimg.bytesPerLine(), 1)
    return to_rgb(img)


def _blank_page(width: int, height: int):
    from PySide6.QtCore import Qt
    from PySide6.QtGui import QImage

    page = QImage(width, height, QImage.Format.Format_RGB32)
    page.fill(Qt.GlobalColor.white)
    dpm = round(96 / 0.0254)
    page.setDotsPerMeterX(dpm)
    page.setDotsPerMeterY(dpm)
    return page


# --------------------------------------------------------------------------- PDF

def load_pdf(path: Path, dpi: int, max_pages: int | None = None) -> list[Image.Image]:
    from PySide6.QtCore import QSize
    from PySide6.QtPdf import QPdfDocument

    # PDFium inside Qt is not thread-safe: two PDFs rendered at once crash the program.
    with _PDF_LOCK:
        doc = QPdfDocument(None)
        doc.load(str(path))
        if doc.status() != QPdfDocument.Status.Ready or doc.pageCount() == 0:
            raise LoadError(f"PDF не открывается ({doc.error().name})")
        rendered = []
        try:
            for i in range(doc.pageCount()):
                if max_pages and i >= max_pages:
                    break
                pts = doc.pagePointSize(i)
                scale = dpi / 72
                w, h = max(1, round(pts.width() * scale)), max(1, round(pts.height() * scale))
                if max(w, h) > MAX_SIDE:
                    k = MAX_SIDE / max(w, h)
                    w, h = round(w * k), round(h * k)
                rendered.append(doc.render(i, QSize(w, h)).copy())
        finally:
            doc.close()
    return [_qimage_to_pil(q) for q in rendered]


# --------------------------------------------------------------------------- SVG

def load_svg(path: Path, dpi: int) -> list[Image.Image]:
    from PySide6.QtCore import QRectF
    from PySide6.QtGui import QPainter
    from PySide6.QtSvg import QSvgRenderer

    renderer = QSvgRenderer(str(path))
    if not renderer.isValid():
        raise LoadError("SVG повреждён или не поддерживается")
    view = renderer.viewBoxF()
    size = renderer.defaultSize()
    base_w = size.width() if size.width() > 0 else view.width()
    base_h = size.height() if size.height() > 0 else view.height()
    if base_w <= 0 or base_h <= 0:
        base_w, base_h = 1000, 1000
    scale = dpi / 96
    w, h = max(1, round(base_w * scale)), max(1, round(base_h * scale))
    if max(w, h) > MAX_SIDE:
        k = MAX_SIDE / max(w, h)
        w, h = round(w * k), round(h * k)
    page = _blank_page(w, h)
    painter = QPainter(page)
    painter.setRenderHint(QPainter.RenderHint.Antialiasing)
    renderer.render(painter, QRectF(0, 0, w, h))
    painter.end()
    return [_qimage_to_pil(page)]


# --------------------------------------------------------------------------- text documents

def _read_text(path: Path) -> str:
    raw = path.read_bytes()
    for enc in ("utf-8-sig", "cp1251", "utf-16"):
        try:
            return raw.decode(enc)
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", errors="replace")


def _render_document(doc, dpi: int, landscape: bool = False, max_pages: int | None = None) -> list[Image.Image]:
    """Paginate a QTextDocument onto A4 pages, like QTextDocument.print() does."""
    from PySide6.QtCore import QRectF, QSizeF
    from PySide6.QtGui import QPainter

    page_w_in, page_h_in = A4_INCHES
    if landscape:
        page_w_in, page_h_in = page_h_in, page_w_in
    # Layout happens in 96-dpi logical units, painting is scaled to the target dpi.
    logical_w, logical_h = page_w_in * 96, page_h_in * 96
    margin = 15 / 25.4 * 96  # 15 mm

    root = doc.rootFrame()
    fmt = root.frameFormat()
    fmt.setMargin(margin)
    root.setFrameFormat(fmt)
    doc.setPageSize(QSizeF(logical_w, logical_h))

    scale = dpi / 96
    px_w, px_h = round(logical_w * scale), round(logical_h * scale)
    pages = []
    for i in range(max(1, doc.pageCount())):
        if max_pages and i >= max_pages:
            break
        page = _blank_page(px_w, px_h)
        painter = QPainter(page)
        painter.setRenderHint(QPainter.RenderHint.Antialiasing)
        painter.setRenderHint(QPainter.RenderHint.TextAntialiasing)
        painter.scale(scale, scale)
        painter.translate(0, -i * logical_h)
        doc.drawContents(painter, QRectF(0, i * logical_h, logical_w, logical_h))
        painter.end()
        pages.append(_qimage_to_pil(page))
    return pages


def load_txt(path: Path, dpi: int, max_pages: int | None = None) -> list[Image.Image]:
    from PySide6.QtGui import QFont, QTextDocument, QTextOption

    doc = QTextDocument()
    font = QFont("Consolas", 10)
    font.setStyleHint(QFont.StyleHint.Monospace)
    doc.setDefaultFont(font)
    opt = doc.defaultTextOption()
    opt.setWrapMode(QTextOption.WrapMode.WrapAtWordBoundaryOrAnywhere)
    doc.setDefaultTextOption(opt)
    text = _read_text(path)
    doc.setPlainText(text if text.strip() else " ")
    return _render_document(doc, dpi, max_pages=max_pages)


def load_html(path: Path, dpi: int, max_pages: int | None = None) -> list[Image.Image]:
    from PySide6.QtCore import QUrl
    from PySide6.QtGui import QFont, QTextDocument

    doc = QTextDocument()
    doc.setDefaultFont(QFont("Segoe UI", 10))
    doc.setBaseUrl(QUrl.fromLocalFile(str(path.parent) + "/"))
    doc.setHtml(_read_text(path))
    return _render_document(doc, dpi, max_pages=max_pages)


def _cell_text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    if hasattr(value, "strftime"):
        return value.strftime("%d.%m.%Y %H:%M").replace(" 00:00", "")
    return str(value)


def load_xlsx(path: Path, dpi: int, max_pages: int | None = None) -> list[Image.Image]:
    try:
        import openpyxl
    except ImportError as exc:
        raise LoadError("нет библиотеки openpyxl") from exc
    from PySide6.QtGui import QFont, QTextDocument

    try:
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:
        raise LoadError(str(exc)) from exc

    pages: list[Image.Image] = []
    try:
        for ws in wb.worksheets:
            rows = []
            for r_idx, row in enumerate(ws.iter_rows(values_only=True)):
                if r_idx >= SHEET_ROW_LIMIT:
                    break
                rows.append([_cell_text(v) for v in row])
            # Trim empty trailing rows/columns that openpyxl reports for formatted cells.
            while rows and not any(rows[-1]):
                rows.pop()
            if not rows:
                continue
            n_cols = max((max((i + 1 for i, v in enumerate(r) if v), default=0) for r in rows), default=0)
            if n_cols == 0:
                continue
            body = []
            for r_idx, row in enumerate(rows):
                tag = "th" if r_idx == 0 else "td"
                cells = "".join(
                    f"<{tag}>{html.escape(row[c]) if c < len(row) else ''}</{tag}>" for c in range(n_cols)
                )
                body.append(f"<tr>{cells}</tr>")
            markup = (
                "<style>table{border-collapse:collapse} th,td{border:1px solid #888;padding:3px 6px}"
                "th{background:#e8e8ef;font-weight:600}</style>"
                f"<h3>{html.escape(ws.title)}</h3>"
                f"<table width='100%' cellspacing='0'>{''.join(body)}</table>"
            )
            doc = QTextDocument()
            doc.setDefaultFont(QFont("Segoe UI", 9))
            doc.setHtml(markup)
            remaining = None if not max_pages else max_pages - len(pages)
            pages.extend(_render_document(doc, dpi, landscape=n_cols > 7, max_pages=remaining))
            if max_pages and len(pages) >= max_pages:
                break
    finally:
        wb.close()
    if not pages:
        raise LoadError("в таблице нет данных")
    return pages


# --------------------------------------------------------------------------- entry point

def load_pages(path: Path, dpi: int = 300, max_pages: int | None = None, preview_size: int | None = None) -> list[Image.Image]:
    """Load ``path`` as a list of RGB pages.

    ``preview_size`` switches to a fast low-resolution mode for previews.
    """
    ext = path.suffix.lower()
    if ext in (".heic", ".heif") and not HEIF_SUPPORTED:
        raise LoadError("нет поддержки HEIC (pillow-heif)")
    if ext in RASTER_EXTS:
        return load_raster(path, max_pages=max_pages, draft_size=preview_size)
    if preview_size:
        dpi = 110
    if ext == ".pdf":
        return load_pdf(path, dpi, max_pages)
    if ext == ".svg":
        return load_svg(path, dpi)
    if ext == ".txt":
        return load_txt(path, dpi, max_pages)
    if ext in (".html", ".htm"):
        return load_html(path, dpi, max_pages)
    if ext in OFFICE_EXTS:
        if ext in (".xlsx", ".xlsm") and not can_convert(path):
            return load_xlsx(path, dpi, max_pages)  # simple table view without an office engine
        try:
            return load_office(path, dpi, max_pages, load_pdf)
        except OfficeError as exc:
            raise LoadError(str(exc)) from exc
    raise LoadError(f"формат {ext} не поддерживается")
